"""Convert the recorded spreadsheet (data/LiDAR_Data_Collection.xlsx, sheet
"Object In Center") into data/formatted_sensor_data.csv, one row per sensor per
test: Test, Sensor, Distance (inches), Angle (degrees)."""
import csv
from pathlib import Path

import openpyxl

DATA_DIR = Path(__file__).resolve().parent.parent / "data"
XLSX_PATH = DATA_DIR / "LiDAR_Data_Collection.xlsx"
CSV_PATH = DATA_DIR / "formatted_sensor_data.csv"

SHEET_NAME = "Object In Center"  # must be this sheet
SENSOR_COLS = [
    (4, 5),
    (9, 10),
    (14, 15),
    (19, 20),
]  # (angleCol, distCol) for sensors 1..4
OFFSETS_DEG = [180.0, 180.0, 0.0, 0.0]  # sensors 1–2: +180°, 3–4: +0°
MM2IN = 1 / 25.4
START_ROW = 4


def convert_excel_to_sensor_format(xlsx):
    sh = openpyxl.load_workbook(xlsx, data_only=True)[SHEET_NAME]
    out, r = [], START_ROW
    while True:
        row_pts, any_data = [], False
        for s_idx, (a_col, d_col) in enumerate(SENSOR_COLS):
            a = sh.cell(row=r, column=a_col).value
            d = sh.cell(row=r, column=d_col).value
            if a is None and d is None:
                continue
            any_data = True
            try:
                a, d = float(a), float(d)
            except (TypeError, ValueError):
                continue
            if d == 65535:
                continue
            ang = ((360.0 - a) + OFFSETS_DEG[s_idx]) % 360.0
            row_pts.append([(d + 10.5) * MM2IN, ang])
        if not any_data:
            break
        if len(row_pts) == 4:
            out.extend(row_pts)
        r += 1
    return out


def write_csv(xlsx=XLSX_PATH, csv_out=CSV_PATH):
    data = convert_excel_to_sensor_format(xlsx)
    print(f"Processed {len(data) // 4} sets")
    with open(csv_out, "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["Test", "Sensor", "Distance (inches)", "Angle (degrees)"])
        for i in range(0, len(data), 4):
            t = i // 4 + 1
            for s, (dist, ang) in enumerate(data[i : i + 4], start=1):
                w.writerow([t, s, dist, ang])


if __name__ == "__main__":
    write_csv()
